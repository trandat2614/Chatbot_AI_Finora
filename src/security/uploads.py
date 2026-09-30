"""Validation for untrusted uploads."""
from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

ALLOWED_CONTENT_TYPES = {
    ".csv": {"text/csv", "application/csv", "text/plain", "application/vnd.ms-excel"},
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/plain", "application/octet-stream"},
    ".pdf": {"application/pdf"},
    ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/zip"},
}


class UnsafeUploadError(ValueError):
    pass


def validate_upload(filename: str, content_type: str | None, contents: bytes, max_bytes: int) -> str:
    original = filename or ""
    safe_name = Path(original).name
    if not safe_name or safe_name != original or "\x00" in original:
        raise UnsafeUploadError("Tên file không hợp lệ.")
    suffix = Path(safe_name).suffix.lower()
    if suffix not in ALLOWED_CONTENT_TYPES:
        raise UnsafeUploadError("Định dạng file không được hỗ trợ.")
    if not contents:
        raise UnsafeUploadError("File đang rỗng.")
    if len(contents) > max_bytes:
        raise UnsafeUploadError("File vượt quá kích thước cho phép.")
    mime = (content_type or "application/octet-stream").split(";", 1)[0].strip().lower()
    if mime not in ALLOWED_CONTENT_TYPES[suffix]:
        raise UnsafeUploadError("MIME type không khớp với định dạng file.")
    if suffix == ".pdf" and not contents.startswith(b"%PDF-"):
        raise UnsafeUploadError("Nội dung PDF không hợp lệ.")
    if suffix == ".xlsx" and not contents.startswith(b"PK\x03\x04"):
        raise UnsafeUploadError("Nội dung XLSX không hợp lệ.")
    if suffix == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(contents)) as archive:
                names = {item.filename.lower() for item in archive.infolist()}
                expanded = sum(item.file_size for item in archive.infolist())
                if expanded > 50 * 1024 * 1024 or expanded > max(len(contents) * 100, 1):
                    raise UnsafeUploadError("XLSX có kích thước giải nén không an toàn.")
                if any("vbaproject" in name or "externallinks/" in name for name in names):
                    raise UnsafeUploadError("XLSX chứa macro hoặc external link không được hỗ trợ.")
        except zipfile.BadZipFile as exc:
            raise UnsafeUploadError("Nội dung XLSX không hợp lệ.") from exc
    if suffix in {".csv", ".txt", ".md"}:
        try:
            contents.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise UnsafeUploadError("File text phải dùng UTF-8.") from exc
    return safe_name


def neutralize_spreadsheet_formula(value: object) -> object:
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def validate_csv_shape(contents: bytes, *, max_rows: int = 200_000, max_columns: int = 200) -> None:
    text = contents.decode("utf-8-sig")
    if "\x00" in text:
        raise UnsafeUploadError("CSV chứa byte NUL không hợp lệ.")
    reader = csv.reader(io.StringIO(text))
    for index, row in enumerate(reader):
        if index > max_rows:
            raise UnsafeUploadError("CSV có quá nhiều dòng.")
        if len(row) > max_columns:
            raise UnsafeUploadError("CSV có quá nhiều cột.")


def validate_xlsx_shape(
    contents: bytes,
    *,
    max_rows: int = 200_000,
    max_columns: int = 200,
    max_sheets: int = 20,
) -> None:
    """Bound workbook shape before pandas materializes every worksheet."""
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        workbook = load_workbook(
            io.BytesIO(contents),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        try:
            if len(workbook.worksheets) > max_sheets:
                raise UnsafeUploadError("XLSX có quá nhiều sheet.")
            total_rows = 0
            for worksheet in workbook.worksheets:
                if worksheet.max_column > max_columns:
                    raise UnsafeUploadError("XLSX có quá nhiều cột.")
                total_rows += worksheet.max_row
                if total_rows > max_rows:
                    raise UnsafeUploadError("XLSX có quá nhiều dòng.")
        finally:
            workbook.close()
    except UnsafeUploadError:
        raise
    except (InvalidFileException, KeyError, OSError, ValueError, zipfile.BadZipFile) as exc:
        raise UnsafeUploadError("Workbook XLSX bị hỏng hoặc không hợp lệ.") from exc
